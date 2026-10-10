import io
import os
from pathlib import Path

import httpx
from dotenv import load_dotenv
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader

from app.schemas import (
    AIDetectRequest,
    AIDetectResponse,
    DocumentExtractionResponse,
    RevisionRequest,
    RevisionResponse,
    SimilarityRequest,
    SimilarityResponse,
    OverlapSpan,
)
from app.ai_detector import analyze_text
from app.similarity import compute_similarity, find_overlaps

load_dotenv(Path(__file__).resolve().parents[1] / ".env")

app = FastAPI(title="TextLens API", version="0.2.0")

_origins = os.getenv("TEXTLENS_CORS_ORIGINS", "*")
_allowed_origins = [origin.strip() for origin in _origins.split(",") if origin.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_allowed_origins or ["*"],
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
)

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARS = 150000


@app.get("/")
def root():
    return {"status": "ok", "service": "TextLens API", "version": "0.2.0"}


@app.get("/health")
def health():
    return {"status": "healthy"}


@app.post("/analyze/ai", response_model=AIDetectResponse)
def analyze_ai(req: AIDetectRequest):
    result = analyze_text(req.text)
    # The detector returns an error dictionary for unsupported lengths/model failures;
    # turn those into clear client errors instead of response-model validation errors.
    if result.get("error"):
        message = str(result["error"])
        status = 503 if "Could not load model" in message else 422
        raise HTTPException(status_code=status, detail=message)
    return AIDetectResponse(**result)


@app.post("/analyze/similarity", response_model=SimilarityResponse)
def analyze_similarity(req: SimilarityRequest):
    score = compute_similarity(req.text_a, req.text_b)
    overlaps = find_overlaps(req.text_a, req.text_b)
    return SimilarityResponse(
        similarity_score=score,
        overlaps=[OverlapSpan(**o) for o in overlaps],
    )


@app.post("/documents/extract", response_model=DocumentExtractionResponse)
async def extract_document(file: UploadFile = File(...)):
    filename = Path(file.filename or "document").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt", ".md"}:
        raise HTTPException(status_code=415, detail="Supported formats are PDF, DOCX, TXT, and MD.")

    data = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Files must be 15 MB or smaller.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(status_code=422, detail="This PDF is password-protected. Upload an unlocked copy.")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            doc = Document(io.BytesIO(data))
            blocks = [p.text for p in doc.paragraphs if p.text.strip()]
            # Include table text because student reports often put content in tables.
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        blocks.append(" | ".join(cells))
            text = "\n\n".join(blocks)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The document could not be parsed. Check the file and try again.") from exc

    text = text.strip()
    if len(text) > MAX_EXTRACTED_CHARS:
        raise HTTPException(status_code=413, detail="Extracted text is too long. Split the document into smaller parts.")
    return DocumentExtractionResponse(text=text, word_count=len(text.split()), filename=filename)


_REVISION_GUIDANCE = {
    "clarity": "Improve clarity, logical flow, transitions, and readability while keeping the author's intent and level of detail.",
    "grammar": "Correct grammar, spelling, punctuation, and awkward phrasing with minimal edits. Preserve the author's voice.",
    "concise": "Make the text more concise by removing redundancy while preserving meaning, necessary detail, and important qualifications.",
    "academic": "Improve precision, structure, and an appropriate academic tone without making the writing needlessly complicated.",
}


@app.post("/writing/revise", response_model=RevisionResponse)
async def revise_writing(req: RevisionRequest):
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise HTTPException(
            status_code=503,
            detail="AI revision is not configured. Set OPENAI_API_KEY on the backend, then restart the API.",
        )

    base_url = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini").strip()
    system_message = (
        "You are a careful writing editor. Revise text for the requested editorial goal. "
        "Preserve the author's intended meaning, technical details, claims, caveats, and existing citations. "
        "Do not invent facts, sources, quotations, data, or references. If a claim or citation appears unsupported, "
        "leave it intact and mention it in a short note instead of fabricating support. Preserve proper nouns and code. "
        "Do not attempt to manipulate or evade AI-writing or plagiarism detectors. Do not add a claim that text is human-written. "
        "Return exactly two sections: REVISED TEXT followed by NOTES. Keep notes brief."
    )
    user_message = f"Editorial goal: {_REVISION_GUIDANCE[req.goal]}\n\nText to revise:\n---\n{req.text}\n---"
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_message},
            {"role": "user", "content": user_message},
        ],
        "temperature": 0.25,
        "max_tokens": 7000,
    }

    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(75.0, connect=10.0)) as client:
            response = await client.post(
                f"{base_url}/chat/completions",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload,
            )
        if response.status_code >= 400:
            # Don't echo provider bodies into logs or responses; they may contain operational detail.
            raise HTTPException(status_code=502, detail="The configured AI provider rejected the revision request. Check the model and API settings.")
        content = response.json()["choices"][0]["message"]["content"]
        if not isinstance(content, str) or not content.strip():
            raise HTTPException(status_code=502, detail="The AI provider returned an empty revision.")
    except HTTPException:
        raise
    except httpx.TimeoutException as exc:
        raise HTTPException(status_code=504, detail="The AI revision request timed out. Try a shorter passage.") from exc
    except (httpx.HTTPError, KeyError, IndexError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=502, detail="The AI provider could not complete the revision. Check your API configuration.") from exc

    import re
    revised = content.strip()
    revised = re.sub(r"(?im)^\s*REVISED TEXT\s*:?\s*\n", "", revised, count=1)
    notes = ""
    parts = re.split(r"(?im)^\s*NOTES\s*:?\s*$", revised, maxsplit=1)
    revised = parts[0].strip()
    if len(parts) > 1:
        notes = parts[1].strip()
    if not revised:
        raise HTTPException(status_code=502, detail="The AI provider returned an unusable revision. Please try again.")
    return RevisionResponse(revised_text=revised, notes=notes)
